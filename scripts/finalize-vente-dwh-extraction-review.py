from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd


REPO_ROOT = Path(__file__).resolve().parents[1]
EXTRACTIONS_ROOT = REPO_ROOT / "Youssef_Extractions"
OUTPUT_DIR = REPO_ROOT / "docs" / "warehouse"
REPORT_JSON = OUTPUT_DIR / "vente_dwh_final_extraction_review.json"
REPORT_MD = OUTPUT_DIR / "VENTE_DWH_FINAL_EXTRACTION_REVIEW.md"

FOLDERS = [
    EXTRACTIONS_ROOT / "1st_Extraction",
    EXTRACTIONS_ROOT / "2nd_Extraction",
    EXTRACTIONS_ROOT / "3rd_Extraction",
    EXTRACTIONS_ROOT / "4th_Extraction",
    EXTRACTIONS_ROOT / "vente_clean_import",
    EXTRACTIONS_ROOT / "vente_bi_enrichment_import",
]


@dataclass
class WorkbookSummary:
    file_name: str
    folder: str
    rows: int
    columns: int
    headers: list[str]
    inferred_role: str
    sample: dict[str, Any]


def clean(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if pd.isna(value) if not isinstance(value, (list, tuple, dict, str)) else False:
        return None
    return value


def normalize(value: Any) -> str:
    return "" if value is None else str(value).strip().upper()


def infer_role(file_name: str, headers: list[str]) -> str:
    name = file_name.upper()
    hs = set(headers)
    table_names = [
        "C_ORDERLINE",
        "C_ORDER",
        "C_INVOICELINE",
        "C_INVOICE",
        "C_BPARTNER_LOCATION",
        "C_BPARTNER",
        "C_BP_GROUP",
        "C_LOCATION",
        "C_CITY",
        "C_REGION",
        "C_PAYMENTTERM",
        "M_PRICELIST",
        "AD_USER",
        "M_PRODUCT_PO",
        "C_BPARTNER_VENDOR",
        "M_PRODUCT_COLLECTION",
        "M_PRODUCT_THEME",
        "M_PRODUCT_TYPE",
        "M_PRODUCT_CATEGORY",
        "M_PRODUCT",
        "C_DOCTYPE",
        "C_TAX",
        "AD_ORG",
        "C_ALLOCATIONLINE",
        "C_ALLOCATIONHDR",
        "C_PAYMENT",
        "M_INOUTLINE",
        "M_INOUT",
        "M_WAREHOUSE",
        "M_LOCATOR",
        "RV_STORAGE",
        "C_SALESREGION",
        "C_COMMISSION",
    ]
    for table in table_names:
        if table in name:
            return table
    if {"C_ORDER_ID", "C_ORDERLINE_ID", "NUM_COMMANDE", "COMMERCIAL", "FOURNISSEUR"}.issubset(hs):
        return "DENORMALIZED_SALES_ORDER_LINE"
    if {"OBJECT_TYPE", "OBJECT_NAME"}.issubset(hs):
        return "SALES_TABLE_DISCOVERY"
    if {"DATE_DEBUT", "DATE_FIN"}.issubset(hs) or "NB_COMMANDES" in hs:
        return "CONTROL_QUERY"
    return "UNKNOWN_OR_LEGACY"


def summarize_xlsx(path: Path) -> list[WorkbookSummary]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summaries: list[WorkbookSummary] = []
    try:
        for sheet in workbook.worksheets:
            max_row = sheet.max_row or 0
            max_col = sheet.max_column or 0
            headers = [normalize(sheet.cell(1, col).value) for col in range(1, max_col + 1)]
            headers = [h for h in headers if h]
            sample: dict[str, Any] = {}
            if max_row >= 2 and headers:
                values = [clean(sheet.cell(2, col).value) for col in range(1, min(max_col, len(headers)) + 1)]
                sample = {headers[i]: values[i] if i < len(values) else None for i in range(min(len(headers), len(values)))}
            summaries.append(
                WorkbookSummary(
                    file_name=path.name,
                    folder=path.parent.name,
                    rows=max(0, max_row - 1) if headers else max_row,
                    columns=len(headers) if headers else max_col,
                    headers=headers,
                    inferred_role=infer_role(path.name, headers),
                    sample=sample,
                )
            )
    finally:
        workbook.close()
    return summaries


def count_csv_rows(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        headers = next(reader, [])
        rows = sum(1 for _ in reader)
    return {"rows": rows, "columns": len(headers), "headers": headers}


def load_csv(name: str, enrichment: bool = False, usecols: list[str] | None = None) -> pd.DataFrame:
    base = EXTRACTIONS_ROOT / ("vente_bi_enrichment_import" if enrichment else "vente_clean_import") / "csv"
    return pd.read_csv(base / f"{name}.csv", dtype="string", usecols=usecols, low_memory=False)


def profile_clean_packages() -> dict[str, Any]:
    clean_dir = EXTRACTIONS_ROOT / "vente_clean_import" / "csv"
    enrich_dir = EXTRACTIONS_ROOT / "vente_bi_enrichment_import" / "csv"
    clean_counts = {p.stem: count_csv_rows(p) for p in sorted(clean_dir.glob("*.csv"))}
    enrich_counts = {p.stem: count_csv_rows(p) for p in sorted(enrich_dir.glob("*.csv"))}

    orders = load_csv("C_ORDER", usecols=["C_ORDER_ID", "DATEORDERED", "C_BPARTNER_ID", "SALESREP_ID", "C_BPARTNER_LOCATION_ID", "C_PAYMENTTERM_ID", "M_PRICELIST_ID", "GRANDTOTAL"])
    invoices = load_csv("C_INVOICE", usecols=["C_INVOICE_ID", "DATEINVOICED", "C_BPARTNER_ID", "SALESREP_ID", "C_BPARTNER_LOCATION_ID", "C_PAYMENTTERM_ID", "M_PRICELIST_ID", "GRANDTOTAL", "ISPAID"])
    order_lines = load_csv("C_ORDERLINE", usecols=["C_ORDERLINE_ID", "C_ORDER_ID", "M_PRODUCT_ID", "LINENETAMT", "QTYORDERED", "QTYDELIVERED", "QTYINVOICED"])
    invoice_lines = load_csv("C_INVOICELINE", usecols=["C_INVOICELINE_ID", "C_INVOICE_ID", "M_PRODUCT_ID", "LINENETAMT", "QTYINVOICED"])
    product_po = load_csv("M_PRODUCT_PO", enrichment=True, usecols=["M_PRODUCT_ID", "C_BPARTNER_ID", "ISCURRENTVENDOR", "ISACTIVE"])
    ad_user = load_csv("AD_USER", enrichment=True, usecols=["AD_USER_ID", "NAME"])
    vendors = load_csv("C_BPARTNER_VENDOR", enrichment=True, usecols=["C_BPARTNER_ID", "NAME"])

    order_dates = pd.to_datetime(orders["DATEORDERED"], errors="coerce")
    invoice_dates = pd.to_datetime(invoices["DATEINVOICED"], errors="coerce")
    order_products = order_lines["M_PRODUCT_ID"].dropna().drop_duplicates()
    invoice_products = invoice_lines["M_PRODUCT_ID"].dropna().drop_duplicates()
    supplier_products = product_po["M_PRODUCT_ID"].dropna().drop_duplicates()
    salesrep_ids = pd.concat([orders["SALESREP_ID"], invoices["SALESREP_ID"]]).dropna().drop_duplicates()
    ad_user_ids = ad_user["AD_USER_ID"].dropna().drop_duplicates()
    vendor_ids = vendors["C_BPARTNER_ID"].dropna().drop_duplicates()

    return {
        "clean_counts": clean_counts,
        "enrichment_counts": enrich_counts,
        "periods": {
            "orders_min": str(order_dates.min()),
            "orders_max": str(order_dates.max()),
            "invoices_min": str(invoice_dates.min()),
            "invoices_max": str(invoice_dates.max()),
        },
        "coverage": {
            "salesrep_ids_total": int(salesrep_ids.nunique()),
            "salesrep_ids_mapped": int(salesrep_ids[salesrep_ids.isin(ad_user_ids)].nunique()),
            "order_products_total": int(order_products.nunique()),
            "order_products_with_supplier": int(order_products[order_products.isin(supplier_products)].nunique()),
            "invoice_products_total": int(invoice_products.nunique()),
            "invoice_products_with_supplier": int(invoice_products[invoice_products.isin(supplier_products)].nunique()),
            "vendor_ids_total": int(product_po["C_BPARTNER_ID"].dropna().nunique()),
            "vendor_ids_mapped": int(product_po["C_BPARTNER_ID"].dropna().drop_duplicates()[product_po["C_BPARTNER_ID"].dropna().drop_duplicates().isin(vendor_ids)].nunique()),
        },
    }


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook_summaries: list[WorkbookSummary] = []
    for folder in FOLDERS:
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*.xlsx"), key=lambda p: p.name.lower()):
            workbook_summaries.extend(summarize_xlsx(path))

    package_profile = profile_clean_packages()
    roles: dict[str, int] = {}
    for summary in workbook_summaries:
        roles[summary.inferred_role] = roles.get(summary.inferred_role, 0) + 1

    required_dw_areas = {
        "orders_and_order_lines": "complete",
        "invoices_and_invoice_lines": "complete",
        "customers": "complete",
        "commercials": "complete",
        "products_categories_types_themes_collections": "complete",
        "suppliers": "complete_enough",
        "delivery_flow": "complete",
        "payments_and_allocations": "complete",
        "geography_locations_regions_cities": "complete",
        "payment_terms": "complete",
        "price_lists": "complete",
        "sales_regions": "complete",
        "commission_objectives": "not_used_or_empty",
        "contracts": "not_found_in_discovery",
    }

    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workbook_summaries": [asdict(item) for item in workbook_summaries],
        "role_counts": roles,
        "package_profile": package_profile,
        "required_dw_areas": required_dw_areas,
        "final_decision": "No more extraction required for the current processus de vente warehouse scope. Move to modeling/import unless the business explicitly asks for longer history or a non-vente module.",
    }
    REPORT_JSON.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    REPORT_MD.write_text(render_markdown(payload), encoding="utf-8")
    print(f"Wrote {REPORT_MD}")
    print(f"Wrote {REPORT_JSON}")


def render_markdown(payload: dict[str, Any]) -> str:
    lines: list[str] = []
    lines.append("# Vente DW Final Extraction Review")
    lines.append("")
    lines.append(f"Generated at: `{payload['generated_at']}`")
    lines.append("")
    lines.append("## Final Decision")
    lines.append("")
    lines.append(payload["final_decision"])
    lines.append("")
    lines.append("## Data Warehouse Area Coverage")
    lines.append("")
    lines.append("| Area | Status |")
    lines.append("|---|---|")
    for area, status in payload["required_dw_areas"].items():
        lines.append(f"| `{area}` | `{status}` |")
    lines.append("")
    lines.append("## Clean Package Coverage")
    lines.append("")
    lines.append("| Table | Rows | Columns |")
    lines.append("|---|---:|---:|")
    for table, meta in sorted(payload["package_profile"]["clean_counts"].items()):
        lines.append(f"| `{table}` | {meta['rows']} | {meta['columns']} |")
    lines.append("")
    lines.append("## Enrichment Package Coverage")
    lines.append("")
    lines.append("| Table | Rows | Columns |")
    lines.append("|---|---:|---:|")
    for table, meta in sorted(payload["package_profile"]["enrichment_counts"].items()):
        lines.append(f"| `{table}` | {meta['rows']} | {meta['columns']} |")
    lines.append("")
    lines.append("## Time Coverage")
    lines.append("")
    for key, value in payload["package_profile"]["periods"].items():
        lines.append(f"- `{key}`: {value}")
    lines.append("")
    lines.append("## Key Mapping Coverage")
    lines.append("")
    for key, value in payload["package_profile"]["coverage"].items():
        lines.append(f"- `{key}`: {value}")
    lines.append("")
    lines.append("## Workbook Classification Counts")
    lines.append("")
    lines.append("| Role | Count |")
    lines.append("|---|---:|")
    for role, count in sorted(payload["role_counts"].items()):
        lines.append(f"| `{role}` | {count} |")
    lines.append("")
    lines.append("## Important Notes")
    lines.append("")
    lines.append("- `3.xlsx` and `4.xlsx` are duplicate denormalized sales/order-line extracts; keep one as a validation/export source, not both.")
    lines.append("- Commission/objective exports are empty or obsolete for this LPN context, so they should not block the vente warehouse.")
    lines.append("- `1.xlsx` and `2.xlsx` are control outputs, useful for traceability but not warehouse tables.")
    lines.append("- For future yearly/seasonality BI, extract a wider historical range. For the current scoped vente warehouse, extraction is sufficient.")
    lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    main()
