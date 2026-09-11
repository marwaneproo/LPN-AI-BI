from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = REPO_ROOT / "Youssef_Extractions" / "4th_Extraction"
OUTPUT_DIR = REPO_ROOT / "docs" / "warehouse"
REPORT_JSON = OUTPUT_DIR / "vente_fourth_extraction_audit.json"
REPORT_MD = OUTPUT_DIR / "VENTE_FOURTH_EXTRACTION_AUDIT.md"


@dataclass
class SheetAudit:
    sheet_name: str
    rows: int
    columns: int
    headers: list[str]
    inferred_content: str
    sample_rows: list[dict[str, Any]]


@dataclass
class WorkbookAudit:
    file_name: str
    path: str
    size_bytes: int
    role: str
    sheets: list[SheetAudit]
    notes: list[str]


def clean(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def normalize(value: Any) -> str:
    return "" if value is None else str(value).strip().upper()


def infer_content(file_name: str, headers: list[str]) -> str:
    name = file_name.upper()
    hs = set(headers)
    if "C_BP_GROUP" in name or "C_BP_GROUP_ID" in hs:
        return "C_BP_GROUP"
    if "C_BPARTNER_LOCATION" in name or "C_BPARTNER_LOCATION_ID" in hs:
        return "C_BPARTNER_LOCATION"
    if "C_LOCATION" in name or "C_LOCATION_ID" in hs:
        return "C_LOCATION"
    if "C_PAYMENTTERM" in name or "C_PAYMENTTERM_ID" in hs:
        return "C_PAYMENTTERM"
    if "SALES_TABLE_DISCOVERY" in name or {"OBJECT_TYPE", "OBJECT_NAME"}.issubset(hs):
        return "SALES_TABLE_DISCOVERY"
    if "C_CONTRACT" in name or "CONTRACT" in hs:
        return "CONTRACT_OR_AGREEMENT_CANDIDATE"
    if "C_COMMISSION" in name or "COMMISSION" in name:
        return "COMMISSION_CANDIDATE"
    if "C_SALES" in name or "SALES" in name:
        return "SALES_CANDIDATE"
    if "C_ORDER" in hs and "GRANDTOTAL" in hs:
        return "ORDER_RELATED"
    return "UNKNOWN"


def role_for_file(file_name: str, sheets: list[SheetAudit]) -> tuple[str, list[str]]:
    notes: list[str] = []
    name = file_name.upper()
    inferred = {sheet.inferred_content for sheet in sheets}
    if name.startswith("27_C_BP_GROUP"):
        return "required_dimension", notes
    if name.startswith("28_C_BPARTNER_LOCATION"):
        return "required_dimension", notes
    if name.startswith("29_C_LOCATION"):
        return "required_dimension", notes
    if name.startswith("30_C_PAYMENTTERM"):
        return "required_dimension", notes
    if name.startswith("31_SALES_TABLE_DISCOVERY"):
        return "metadata_discovery", notes
    if file_name in {"1.xlsx", "2.xlsx", "3.xlsx", "4.xlsx"}:
        notes.append("Additional ad-hoc extraction; classification depends on headers/content.")
        return "additional_candidate", notes
    if "UNKNOWN" in inferred:
        notes.append("Could not confidently classify from filename/header.")
    return "raw_candidate", notes


def read_workbook(path: Path) -> list[SheetAudit]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        result: list[SheetAudit] = []
        for ws in wb.worksheets:
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            headers = [normalize(ws.cell(1, col).value) for col in range(1, max_col + 1)]
            headers = [header for header in headers if header]
            samples: list[dict[str, Any]] = []
            for row_idx in range(2, min(max_row, 6) + 1):
                values = [clean(ws.cell(row_idx, col).value) for col in range(1, min(max_col, len(headers)) + 1)]
                if any(value is not None for value in values):
                    samples.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
            rows = max(0, max_row - 1) if headers else max_row
            result.append(
                SheetAudit(
                    sheet_name=ws.title,
                    rows=rows,
                    columns=len(headers) if headers else max_col,
                    headers=headers,
                    inferred_content=infer_content(path.name, headers),
                    sample_rows=samples,
                )
            )
        return result
    finally:
        wb.close()


def summarize_discovery(audits: list[WorkbookAudit]) -> dict[str, list[str]]:
    terms = ("CONTRACT", "AGREEMENT", "COMMISSION", "SALES", "TARGET", "OBJECTIVE", "GOAL")
    grouped: dict[str, list[str]] = defaultdict(list)
    for audit in audits:
        if "DISCOVERY" not in audit.file_name.upper():
            continue
        for sheet in audit.sheets:
            if not {"OBJECT_TYPE", "OBJECT_NAME"}.issubset(set(sheet.headers)):
                continue
            for row in sheet.sample_rows:
                pass
            path = Path(audit.path)
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[sheet.sheet_name]
                header_to_idx = {normalize(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}
                object_name_col = header_to_idx.get("OBJECT_NAME")
                if not object_name_col:
                    continue
                for row_idx in range(2, (ws.max_row or 1) + 1):
                    object_name = normalize(ws.cell(row_idx, object_name_col).value)
                    if not object_name:
                        continue
                    matched = False
                    for term in terms:
                        if term in object_name:
                            grouped[term].append(object_name)
                            matched = True
                    if not matched:
                        grouped["OTHER"].append(object_name)
            finally:
                wb.close()
    return {key: sorted(set(values)) for key, values in sorted(grouped.items())}


def build_recommendations(audits: list[WorkbookAudit], discovery: dict[str, list[str]]) -> list[str]:
    contents = {sheet.inferred_content for audit in audits for sheet in audit.sheets}
    recommendations: list[str] = []
    required = ["C_BP_GROUP", "C_BPARTNER_LOCATION", "C_LOCATION", "C_PAYMENTTERM"]
    missing = [item for item in required if item not in contents]
    if missing:
        recommendations.append(f"Missing required dimensions still not validated: {', '.join(missing)}.")
    if discovery:
        recommendations.append("Review discovery objects before extracting contract/objective/commission facts.")
    if not any(discovery.get(term) for term in ("CONTRACT", "AGREEMENT", "COMMISSION", "TARGET", "OBJECTIVE", "GOAL")):
        recommendations.append("Discovery did not expose obvious contract/objective/commission tables in the sampled names.")
    recommendations.append("Current 4th extraction closes customer group, geography, and payment-term dimensions if row counts/headers are valid.")
    return recommendations


def build_markdown(audits: list[WorkbookAudit], discovery: dict[str, list[str]], recommendations: list[str]) -> str:
    lines: list[str] = []
    lines.append("# Vente Fourth Extraction Audit")
    lines.append("")
    lines.append(f"Generated at: `{datetime.now(timezone.utc).isoformat()}`")
    lines.append("")
    lines.append("## Workbook Inventory")
    lines.append("")
    lines.append("| File | Role | Sheet | Rows | Columns | Inferred content | Main headers |")
    lines.append("|---|---|---|---:|---:|---|---|")
    for audit in audits:
        for sheet in audit.sheets:
            headers = ", ".join(sheet.headers[:10])
            lines.append(
                f"| `{audit.file_name}` | `{audit.role}` | `{sheet.sheet_name}` | {sheet.rows} | {sheet.columns} | "
                f"`{sheet.inferred_content}` | {headers} |"
            )
    lines.append("")
    lines.append("## Additional Candidate Details")
    lines.append("")
    for audit in audits:
        if audit.role != "additional_candidate":
            continue
        lines.append(f"### `{audit.file_name}`")
        lines.append("")
        for sheet in audit.sheets:
            lines.append(f"- Sheet `{sheet.sheet_name}`: `{sheet.inferred_content}`, {sheet.rows} rows, {sheet.columns} columns.")
            if sheet.sample_rows:
                compact = {k: v for k, v in list(sheet.sample_rows[0].items())[:8]}
                lines.append(f"- First sample: `{json.dumps(compact, ensure_ascii=False)}`")
        lines.append("")
    lines.append("## Discovery Summary")
    lines.append("")
    if discovery:
        for term, names in discovery.items():
            lines.append(f"### `{term}`")
            lines.append("")
            for name in names[:80]:
                lines.append(f"- `{name}`")
            if len(names) > 80:
                lines.append(f"- ... {len(names) - 80} more")
            lines.append("")
    else:
        lines.append("No discovery objects found or readable.")
        lines.append("")
    lines.append("## Recommendations")
    lines.append("")
    for recommendation in recommendations:
        lines.append(f"- {recommendation}")
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    audits: list[WorkbookAudit] = []
    for path in sorted(SOURCE_DIR.glob("*.xlsx"), key=lambda item: item.name.lower()):
        sheets = read_workbook(path)
        role, notes = role_for_file(path.name, sheets)
        audits.append(
            WorkbookAudit(
                file_name=path.name,
                path=str(path),
                size_bytes=path.stat().st_size,
                role=role,
                sheets=sheets,
                notes=notes,
            )
        )

    discovery = summarize_discovery(audits)
    recommendations = build_recommendations(audits, discovery)
    REPORT_JSON.write_text(
        json.dumps(
            {
                "generated_at": datetime.now(timezone.utc).isoformat(),
                "source_dir": str(SOURCE_DIR),
                "workbooks": [asdict(audit) for audit in audits],
                "discovery_summary": discovery,
                "recommendations": recommendations,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    REPORT_MD.write_text(build_markdown(audits, discovery, recommendations), encoding="utf-8")
    print(f"Wrote {REPORT_MD}")
    print(f"Wrote {REPORT_JSON}")
    print(f"Audited {len(audits)} workbooks")


if __name__ == "__main__":
    main()
