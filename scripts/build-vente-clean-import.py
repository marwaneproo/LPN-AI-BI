from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from dataclasses import dataclass
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

import openpyxl


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE_DIR = REPO_ROOT / "Youssef_Extractions" / "2nd_Extraction"
DEFAULT_OUTPUT_DIR = REPO_ROOT / "Youssef_Extractions" / "vente_clean_import"


@dataclass(frozen=True)
class TableSource:
    table_name: str
    source_file: str
    sheet_name: str
    required_headers: tuple[str, ...]
    filter_used: str
    notes: str = ""


VENTE_FILTER = (
    "Vente 4-month extraction. Orders: DATEORDERED >= ADD_MONTHS(TRUNC(SYSDATE), -4), "
    "ISSOTRX='Y', DOCSTATUS IN ('CO','CL'), C_DOCTYPETARGET_ID IN "
    "(1000028,1000034,1000032,1000031). Invoices: DATEINVOICED >= "
    "ADD_MONTHS(TRUNC(SYSDATE), -4), ISSOTRX='Y', DOCSTATUS='CO', "
    "C_DOCTYPE_ID IN (1000002,1000003,1000004)."
)


TABLE_SOURCES: tuple[TableSource, ...] = (
    TableSource(
        table_name="C_ORDER",
        source_file="Book5.xlsx",
        sheet_name="Sheet1",
        required_headers=("C_ORDER_ID", "DOCUMENTNO", "DOCSTATUS", "ISSOTRX"),
        filter_used=VENTE_FILTER,
        notes="Canonical order header export. Duplicate order headers inside 02_C_ORDERLINE workbooks are ignored.",
    ),
    TableSource(
        table_name="C_ORDERLINE",
        source_file="02_C_ORDERLINE.xlsx",
        sheet_name="Sheet2",
        required_headers=("C_ORDERLINE_ID", "C_ORDER_ID", "M_PRODUCT_ID", "QTYORDERED"),
        filter_used=VENTE_FILTER,
    ),
    TableSource(
        table_name="C_INVOICE",
        source_file="03_C_INVOICE.xlsx",
        sheet_name="Sheet1",
        required_headers=("C_INVOICE_ID", "DOCUMENTNO", "DOCSTATUS", "ISSOTRX"),
        filter_used=VENTE_FILTER,
    ),
    TableSource(
        table_name="C_INVOICELINE",
        source_file="04_C_INVOICELINE.xlsx",
        sheet_name="Sheet2",
        required_headers=("C_INVOICELINE_ID", "C_INVOICE_ID", "M_PRODUCT_ID", "LINENETAMT"),
        filter_used=VENTE_FILTER,
    ),
    TableSource(
        table_name="C_BPARTNER",
        source_file="05_C_BPARTNER.xlsx",
        sheet_name="Sheet1",
        required_headers=("C_BPARTNER_ID", "VALUE", "NAME"),
        filter_used="Business partners referenced by the 4-month vente order and invoice extracts.",
    ),
    TableSource(
        table_name="M_PRODUCT",
        source_file="06_M_PRODUCT.xlsx",
        sheet_name="Sheet2",
        required_headers=("M_PRODUCT_ID", "VALUE", "NAME", "M_PRODUCT_CATEGORY_ID"),
        filter_used="Products referenced by the 4-month vente order-line and invoice-line extracts.",
    ),
    TableSource(
        table_name="M_PRODUCT_CATEGORY",
        source_file="07_M_PRODUCT_CATEGORY.xlsx",
        sheet_name="Sheet1",
        required_headers=("M_PRODUCT_CATEGORY_ID", "VALUE", "NAME"),
        filter_used="Product categories referenced by the extracted vente products.",
    ),
    TableSource(
        table_name="C_DOCTYPE",
        source_file="08_C_DOCTYPE.xlsx",
        sheet_name="Sheet1",
        required_headers=("C_DOCTYPE_ID", "NAME", "DOCBASETYPE"),
        filter_used="Document types referenced by vente orders and invoices.",
    ),
    TableSource(
        table_name="C_TAX",
        source_file="09_C_TAX.xlsx",
        sheet_name="Sheet2",
        required_headers=("C_TAX_ID", "NAME", "RATE"),
        filter_used="Taxes referenced by extracted vente order lines and invoice lines.",
    ),
    TableSource(
        table_name="AD_ORG",
        source_file="10_AD_ORG.xlsx",
        sheet_name="Sheet3",
        required_headers=("AD_ORG_ID", "VALUE", "NAME"),
        filter_used="Organizations referenced by extracted vente records.",
    ),
    TableSource(
        table_name="C_ALLOCATIONLINE",
        source_file="11_C_ALLOCATIONLINE.xlsx",
        sheet_name="Sheet2",
        required_headers=("C_ALLOCATIONLINE_ID", "C_INVOICE_ID", "C_PAYMENT_ID", "AMOUNT"),
        filter_used="Allocation lines linked to extracted vente invoices.",
    ),
    TableSource(
        table_name="C_ALLOCATIONHDR",
        source_file="12_C_ALLOCATIONHDR.xlsx",
        sheet_name="Sheet1",
        required_headers=("C_ALLOCATIONHDR_ID", "DOCUMENTNO", "DATETRX", "DOCSTATUS"),
        filter_used="Allocation headers linked to extracted vente allocations.",
    ),
    TableSource(
        table_name="C_PAYMENT",
        source_file="13_C_PAYMENT.xlsx",
        sheet_name="Sheet2",
        required_headers=("C_PAYMENT_ID", "DOCUMENTNO", "DATETRX", "PAYAMT"),
        filter_used="Payments linked through C_ALLOCATIONLINE.C_PAYMENT_ID for extracted vente invoices.",
        notes="C_PAYMENT.C_INVOICE_ID is empty in this LPN database; use allocations for invoice-payment joins.",
    ),
    TableSource(
        table_name="M_INOUT",
        source_file="14_M_INOUT.xlsx",
        sheet_name="Sheet3",
        required_headers=("M_INOUT_ID", "DOCUMENTNO", "C_ORDER_ID", "MOVEMENTDATE"),
        filter_used="Shipments/deliveries linked to extracted vente orders.",
    ),
    TableSource(
        table_name="M_INOUTLINE",
        source_file="15_M_INOUTLINE.xlsx",
        sheet_name="Sheet4",
        required_headers=("M_INOUTLINE_ID", "M_INOUT_ID", "C_ORDERLINE_ID", "MOVEMENTQTY"),
        filter_used="Shipment/delivery lines linked to extracted vente shipments.",
    ),
    TableSource(
        table_name="M_WAREHOUSE",
        source_file="16_M_WAREHOUSE.xlsx",
        sheet_name="Sheet5",
        required_headers=("M_WAREHOUSE_ID", "VALUE", "NAME"),
        filter_used="Warehouses referenced by extracted vente shipments.",
    ),
    TableSource(
        table_name="M_LOCATOR",
        source_file="17_M_LOCATOR.xlsx",
        sheet_name="Sheet6",
        required_headers=("M_LOCATOR_ID", "VALUE", "M_WAREHOUSE_ID"),
        filter_used="Locators referenced by extracted vente shipment lines.",
    ),
    TableSource(
        table_name="RV_STORAGE",
        source_file="18_RV_STORAGE_SALES_PRODUCTS.xlsx",
        sheet_name="Sheet1",
        required_headers=("M_PRODUCT_ID", "VALUE", "NAME", "QTYONHAND", "QTYRESERVED", "QTYAVAILABLE"),
        filter_used="RV_STORAGE rows for products referenced by extracted vente orders and invoices.",
        notes="Chosen over M_STORAGE because it includes product labels and QTYAVAILABLE.",
    ),
)

DIAGNOSTIC_SELECTED_COUNT_KEYS = {
    "C_ORDER": "C_ORDER_SELECTED",
    "C_ORDERLINE": "C_ORDERLINE_SELECTED",
    "C_INVOICE": "C_INVOICE_SELECTED",
    "C_INVOICELINE": "C_INVOICELINE_SELECTED",
    "C_BPARTNER": "C_BPARTNER_SELECTED",
    "M_PRODUCT": "M_PRODUCT_SELECTED",
}


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build a clean CSV staging package from Youssef's second vente Excel extraction."
    )
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--keep-existing", action="store_true", help="Do not clear the output directory first.")
    args = parser.parse_args()

    source_dir = args.source_dir.resolve()
    output_dir = args.output_dir.resolve()
    _validate_paths(source_dir, output_dir)

    if output_dir.exists() and not args.keep_existing:
        shutil.rmtree(output_dir)

    csv_dir = output_dir / "csv"
    csv_dir.mkdir(parents=True, exist_ok=True)

    manifest_tables: list[dict[str, Any]] = []
    source_mappings: list[dict[str, Any]] = []
    diagnostic_counts = _load_diagnostic_selected_counts(source_dir)

    for source in TABLE_SOURCES:
        source_path = source_dir / source.source_file
        csv_path = csv_dir / f"{source.table_name}.csv"
        result = _convert_sheet_to_csv(source, source_path, csv_path)

        manifest_tables.append(
            {
                "name": source.table_name,
                "csv_file": f"csv/{source.table_name}.csv",
                "row_count": result["row_count"],
                "exported_columns": result["headers"],
                "filter_used": f"{source.filter_used} Source: {source.source_file}::{source.sheet_name}.",
            }
        )
        source_mappings.append(
            {
                "table_name": source.table_name,
                "csv_file": f"csv/{source.table_name}.csv",
                "source_file": source.source_file,
                "source_sheet": source.sheet_name,
                "source_sha256": _sha256(source_path),
                "source_excel_rows": result["source_excel_rows"],
                "source_excel_columns": result["source_excel_columns"],
                "csv_data_rows": result["row_count"],
                "csv_columns": len(result["headers"]),
                "diagnostic_expected_rows": diagnostic_counts.get(source.table_name),
                "diagnostic_delta": (
                    result["row_count"] - diagnostic_counts[source.table_name]
                    if source.table_name in diagnostic_counts
                    else None
                ),
                "required_headers_checked": list(source.required_headers),
                "notes": source.notes,
            }
        )

    ignored_files = _ignored_files(source_dir)
    exported_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    manifest = {
        "snapshot_id": "lpn-vente-4months-20260514-v1",
        "exported_at": exported_at,
        "exported_by": "Youssef Bahaddou",
        "source_system": "Compiere on Oracle 11g - LPN Mohammedia",
        "tables": manifest_tables,
        "notes": (
            "Clean vente-only CSV staging package generated from "
            "Youssef_Extractions/2nd_Extraction. Diagnostics and duplicate Excel exports are excluded."
        ),
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    (output_dir / "source_mapping.json").write_text(
        json.dumps(
            {
                "generated_at": exported_at,
                "source_dir": str(source_dir),
                "output_dir": str(output_dir),
                "tables": source_mappings,
                "ignored_files": ignored_files,
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    _write_report(output_dir, manifest_tables, source_mappings, ignored_files)

    print(f"Created clean vente staging package: {output_dir}")
    print(f"CSV files: {csv_dir}")
    print(f"Tables: {len(manifest_tables)}")
    print(f"Rows: {sum(table['row_count'] for table in manifest_tables)}")


def _validate_paths(source_dir: Path, output_dir: Path) -> None:
    if not source_dir.exists():
        raise SystemExit(f"Source directory does not exist: {source_dir}")
    if not source_dir.is_dir():
        raise SystemExit(f"Source path is not a directory: {source_dir}")

    expected_parent = (REPO_ROOT / "Youssef_Extractions").resolve()
    try:
        output_dir.relative_to(expected_parent)
    except ValueError as exc:
        raise SystemExit(
            "Refusing to write outside Youssef_Extractions. "
            f"Output was: {output_dir}"
        ) from exc

    missing = [source.source_file for source in TABLE_SOURCES if not (source_dir / source.source_file).exists()]
    if missing:
        raise SystemExit(f"Missing required source files: {', '.join(missing)}")


def _convert_sheet_to_csv(source: TableSource, source_path: Path, csv_path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
    if source.sheet_name not in workbook.sheetnames:
        raise SystemExit(f"{source.source_file} does not contain required sheet {source.sheet_name}")

    worksheet = workbook[source.sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise SystemExit(f"{source.source_file}::{source.sheet_name} is empty") from exc

    headers = _normalize_headers(raw_headers)
    _validate_headers(source, headers)

    row_count = 0
    csv_path.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(headers)
        for row in rows:
            values = list(row[: len(headers)])
            if _row_is_empty(values):
                continue
            writer.writerow([_clean_cell(value) for value in values])
            row_count += 1

    return {
        "headers": headers,
        "row_count": row_count,
        "source_excel_rows": worksheet.max_row,
        "source_excel_columns": worksheet.max_column,
    }


def _normalize_headers(raw_headers: tuple[Any, ...]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw_headers, start=1):
        header = str(value).strip().upper() if value not in (None, "") else f"UNNAMED_{index}"
        header = " ".join(header.split())
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 1
        headers.append(header)
    while headers and headers[-1].startswith("UNNAMED_"):
        headers.pop()
    return headers


def _validate_headers(source: TableSource, headers: list[str]) -> None:
    missing = [header for header in source.required_headers if header not in headers]
    if missing:
        raise SystemExit(
            f"{source.source_file}::{source.sheet_name} does not look like {source.table_name}; "
            f"missing required headers: {', '.join(missing)}"
        )


def _row_is_empty(values: list[Any]) -> bool:
    return all(value is None or str(value).strip() == "" for value in values)


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, bool):
        return "Y" if value else "N"
    text = str(value)
    text = text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    text = " ".join(text.split()) if "\t" in text else text.strip()
    return text


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _ignored_files(source_dir: Path) -> list[dict[str, str]]:
    used = {source.source_file for source in TABLE_SOURCES}
    ignored: list[dict[str, str]] = []
    for path in sorted(source_dir.iterdir(), key=lambda item: item.name.lower()):
        if path.name in used:
            continue
        if path.suffix.lower() != ".xlsx":
            reason = "not a workbook data export"
        elif path.name.startswith("02_C_ORDERLINE("):
            reason = "duplicate of canonical 02_C_ORDERLINE.xlsx"
        elif path.name.startswith("17_M_LOCATOR("):
            reason = "duplicate of canonical 17_M_LOCATOR.xlsx"
        elif path.name.startswith("Book"):
            reason = "diagnostic/discovery workbook, not a business table import"
        elif path.name.startswith("18A_") or path.name.startswith("18B_"):
            reason = "column discovery workbook, not a business table import"
        else:
            reason = "not selected for vente clean import"
        ignored.append({"file": path.name, "reason": reason})
    return ignored


def _load_diagnostic_selected_counts(source_dir: Path) -> dict[str, int]:
    path = source_dir / "Book4.xlsx"
    if not path.exists():
        return {}

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    raw_counts: dict[str, int] = {}
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        try:
            raw_counts[str(row[0])] = int(row[1])
        except (TypeError, ValueError):
            continue

    counts: dict[str, int] = {}
    for table_name, diagnostic_key in DIAGNOSTIC_SELECTED_COUNT_KEYS.items():
        if diagnostic_key in raw_counts:
            counts[table_name] = raw_counts[diagnostic_key]
    return counts


def _write_report(
    output_dir: Path,
    manifest_tables: list[dict[str, Any]],
    source_mappings: list[dict[str, Any]],
    ignored_files: list[dict[str, str]],
) -> None:
    lines = [
        "# Vente Clean Import Report",
        "",
        "Generated by `scripts/build-vente-clean-import.py`.",
        "",
        "## Tables",
        "",
        "| Table | Rows | Diagnostic rows | Delta | Columns | Source | Sheet |",
        "|---|---:|---:|---:|---:|---|---|",
    ]
    by_table = {mapping["table_name"]: mapping for mapping in source_mappings}
    for table in manifest_tables:
        mapping = by_table[table["name"]]
        expected = mapping["diagnostic_expected_rows"]
        delta = mapping["diagnostic_delta"]
        lines.append(
            f"| `{table['name']}` | {table['row_count']} | "
            f"{expected if expected is not None else ''} | "
            f"{delta if delta is not None else ''} | "
            f"{len(table['exported_columns'])} | "
            f"`{mapping['source_file']}` | `{mapping['source_sheet']}` |"
        )

    lines.extend(["", "## Ignored Files", ""])
    for ignored in ignored_files:
        lines.append(f"- `{ignored['file']}`: {ignored['reason']}")

    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- CSVs are encoded as UTF-8 with BOM for Excel friendliness and Python import compatibility.",
            "- Diagnostic workbooks and duplicate exports are intentionally excluded.",
            "- Some Oracle-exported text contains encoding artifacts from the source Excel export; values are preserved.",
        ]
    )
    (output_dir / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
