from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\LPN_PROJECT")
EXPORT_DIR = ROOT / "Youssef_Extractions" / "data" / "Exported_data_through_a_drive"
REPORT_PATH = ROOT / "Youssef_Extractions" / "data" / "drive_export_package_profile.json"

EXPECTED_NEXT_EXPORTS = [
    "59_COMMERCIAL_CUSTOMER_PORTFOLIO_24M_FIXED.xlsx",
    "60_M_INOUT_HEADER_24M.xlsx",
    "61_M_INOUT_LINE_24M.xlsx",
    "62_C_PAYMENT_DETAIL_24M.xlsx",
    "63_C_ALLOCATION_DETAIL_24M.xlsx",
]

KEY_FILES = [
    "53_COMMERCIAL_ORDER_HEADER_24M.xlsx",
    "54_COMMERCIAL_ORDER_LINE_24M.xlsx",
    "55_COMMERCIAL_INVOICE_HEADER_24M.xlsx",
    "56_COMMERCIAL_INVOICE_LINE_24M.xlsx",
    "57_COMMERCIAL_CUSTOMER_PORTFOLIO_24M.xlsx",
    "58_COMMERCIAL_CUSTOMER_LOCATION_24M.xlsx",
]


def workbook_profile(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            headers = []
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                headers = [str(value).strip() if value is not None else "" for value in row]
                break
            sheets.append(
                {
                    "sheet": ws.title,
                    "max_row": ws.max_row,
                    "data_rows": max(ws.max_row - 1, 0),
                    "max_column": ws.max_column,
                    "columns": headers,
                }
            )
    finally:
        wb.close()

    return {
        "name": path.name,
        "size_bytes": path.stat().st_size,
        "sheets": sheets,
    }


def main() -> None:
    if not EXPORT_DIR.exists():
        raise SystemExit(f"Export directory does not exist: {EXPORT_DIR}")

    files = sorted(EXPORT_DIR.glob("*.xlsx"))
    existing_names = {path.name for path in files}

    all_files = [
        {
            "name": path.name,
            "size_bytes": path.stat().st_size,
            "modified": path.stat().st_mtime,
        }
        for path in files
    ]

    key_profiles = []
    for name in KEY_FILES:
        path = EXPORT_DIR / name
        if path.exists():
            key_profiles.append(workbook_profile(path))

    report = {
        "export_dir": str(EXPORT_DIR),
        "file_count": len(files),
        "total_size_bytes": sum(path.stat().st_size for path in files),
        "expected_next_exports_present": [
            name for name in EXPECTED_NEXT_EXPORTS if name in existing_names
        ],
        "expected_next_exports_missing": [
            name for name in EXPECTED_NEXT_EXPORTS if name not in existing_names
        ],
        "commercial_24m_key_files_present": [
            name for name in KEY_FILES if name in existing_names
        ],
        "all_files": all_files,
        "key_profiles": key_profiles,
    }

    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"Export directory: {EXPORT_DIR}")
    print(f"XLSX files: {report['file_count']}")
    print(f"Total size MB: {report['total_size_bytes'] / 1024 / 1024:.2f}")
    print("Expected next exports present:", ", ".join(report["expected_next_exports_present"]) or "none")
    print("Expected next exports missing:", ", ".join(report["expected_next_exports_missing"]) or "none")
    print("Commercial 24M key files present:", ", ".join(report["commercial_24m_key_files_present"]) or "none")
    print()
    for profile in key_profiles:
        sheet = profile["sheets"][0]
        print(
            f"{profile['name']}: {sheet['data_rows']} rows, "
            f"{sheet['max_column']} columns, {profile['size_bytes'] / 1024 / 1024:.2f} MB"
        )
        print("  Columns:", ", ".join(sheet["columns"][:12]), "...")
    print()
    print(f"Wrote report: {REPORT_PATH}")


if __name__ == "__main__":
    main()
